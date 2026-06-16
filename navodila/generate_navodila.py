#!/usr/bin/env python3
"""
Generator navodil za BIO Alinity.

Prebere NAVODILA.md (omejen Markdown) in posnetke v screens/ ter sestavi
Word datoteko Navodila_BioAlinity.docx. Različico aplikacije prebere iz
okoljske spremenljivke APP_BUILD (v CI = github.run_number; versionName je
"1.0.<run> (<sha>)"), da se navodila ujemajo z aplikacijo.

Zagon:  python generate_navodila.py
Odvisnost:  python-docx  (pip install python-docx)
"""
import os
import re
import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
MD = HERE / "NAVODILA.md"
SCREENS = HERE / "screens"
OUT = HERE / "Navodila_BioAlinity.docx"

IMG_WIDTH = Inches(6.2)  # ležeči posnetki tablice (široki)


def app_version():
    """Različica iz CI: versionName = '1.0.<APP_BUILD> (<sha>)'."""
    build = os.environ.get("APP_BUILD") or os.environ.get("GITHUB_RUN_NUMBER")
    sha = os.environ.get("APP_SHA") or os.environ.get("GITHUB_SHA")
    if build:
        v = "v1.0." + str(build)
        if sha:
            v += " (" + str(sha)[:7] + ")"
        return v
    return None


def add_runs(paragraph, text):
    """Doda besedilo z osnovnim **krepkim** oblikovanjem."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part == "":
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:  # liho = znotraj **...**
            run.bold = True


def add_image(doc, alt, rel):
    img = (HERE / rel).resolve()
    if not img.exists():
        doc.add_paragraph(f"[manjka posnetek: {rel}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img), width=IMG_WIDTH)
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(alt)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def build():
    doc = Document()
    # privzeta pisava
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = MD.read_text(encoding="utf-8").splitlines()
    img_re = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
    first_h1 = True

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        m = img_re.match(line)
        if m:
            add_image(doc, m.group(1), m.group(2))
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            if first_h1:
                doc.add_heading(line[2:], level=0)
                first_h1 = False
                # podnaslov z različico in datumom
                v = app_version()
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = sub.add_run(
                    ("Različica aplikacije: %s   ·   " % v if v else "")
                    + "Navodila posodobljena: "
                    + datetime.date.today().strftime("%d.%m.%Y")
                )
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                doc.add_heading(line[2:], level=1)
        elif re.match(r"^\d+\.\s+", line):
            txt = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_runs(p, txt)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, line)

    doc.save(str(OUT))
    print("Ustvarjeno:", OUT)


if __name__ == "__main__":
    build()
